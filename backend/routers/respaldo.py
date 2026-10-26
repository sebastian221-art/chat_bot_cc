"""
📄 ARCHIVO: backend/routers/respaldo.py

Respaldo completo en Excel: TODO lo que tiene el bot y el panel, en un solo archivo.

  GET /respaldo/excel   (requiere sesión del panel)

Hojas del archivo:
  1. "Léeme"            qué trae el archivo, fecha y cuántos registros tiene cada tabla.
  2. "Revisión"         lo que falta o está mal: se revisa en código cada local, evento, sorteo,
                        promoción, entrada de conocimiento y dato general (sin IA, siempre igual).
  3. "Así ve Any los locales"  cada local tal como lo usa Any 2.0 (zona, rubro, qué vende, WhatsApp,
                        carta, Instagram, fotos).
  4. "Any · …"          configuración de Any 2.0: mensajes fijos, sinónimos, reglas, prompts y modelos.
  5. Una hoja por cada tabla de la base de datos, con TODAS sus columnas y filas.
     (Las contraseñas de los usuarios del panel no se exportan.)

El archivo trae datos personales de clientes (teléfonos y conversaciones): guárdalo con cuidado.
"""
from __future__ import annotations

import io
import json
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

from models.database import get_db
from routers.auth import get_current_user

router = APIRouter(prefix="/respaldo", tags=["respaldo"], dependencies=[Depends(get_current_user)])

COLUMNAS_SECRETAS = {"hashed_password", "password", "password_hash", "token", "secret"}
MAX_CELDA = 32000          # límite de Excel por celda: 32.767 caracteres
MAX_FILAS = 1_000_000      # límite de Excel por hoja: 1.048.576 filas

ZONAS_OFICIALES = ["Pasillo Barichara", "Pasillo Vélez", "Pasillo Socorro", "Pasillo Charalá", "Pasillo Páramo",
                   "Pasillo Curití", "Plaza Principal", "Zona Terrazas", "Plazoleta de Comidas La Ceiba",
                   "Paseo de la Cruz", "Paseo San Gil", "Paseo Chicamocha", "Dunas Club Bar", "Torre Médica",
                   "Oficinas Administrativas"]


# ── utilidades ────────────────────────────────────────────────────

def _celda(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d %H:%M") if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, (bytes, bytearray)):
        return f"(binario, {len(v)} bytes)"
    if isinstance(v, (dict, list)):
        v = json.dumps(v, ensure_ascii=False)
    if isinstance(v, bool):
        return "sí" if v else "no"
    if isinstance(v, (int, float)):
        return v
    s = str(v)
    s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)          # caracteres que Excel no acepta
    return s if len(s) <= MAX_CELDA else s[:MAX_CELDA] + " …(cortado)"


def _nombre_hoja(nombre: str, usados: set) -> str:
    n = re.sub(r"[\[\]\*\?/\\:]", " ", nombre)[:31].strip() or "hoja"
    base, i = n, 2
    while n.lower() in usados:
        sufijo = f" ({i})"
        n = base[:31 - len(sufijo)] + sufijo
        i += 1
    usados.add(n.lower())
    return n


def _norm(t: str) -> str:
    import unicodedata
    t = unicodedata.normalize("NFD", (t or "").lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def _gritado(t: str) -> bool:
    letras = [c for c in (t or "") if c.isalpha()]
    return len(letras) >= 12 and sum(c.isupper() for c in letras) / len(letras) > 0.6


def _vencido(texto: str | None, hoy: date) -> bool | None:
    try:
        from any.conocimiento.vigencia import vigente
        if not texto:
            return None
        return not vigente(texto, hoy)
    except Exception:
        return None


# ── revisión de calidad de los datos ──────────────────────────────

class Revision:
    def __init__(self):
        self.filas: list[list] = []

    def add(self, gravedad: str, area: str, registro: str, problema: str, que_hacer: str, id_=""):
        self.filas.append([gravedad, area, registro, problema, que_hacer, id_])


def revisar(db: Session) -> list[list]:
    rv = Revision()
    hoy = date.today()
    try:
        from any.conocimiento.servicios import info_mall
        info = info_mall(db)
    except Exception:
        info = {}

    # Datos generales del centro comercial
    for clave, etiqueta in [("direccion", "Dirección"), ("horario_general", "Horario general"),
                            ("telefono", "Teléfono"), ("parqueadero", "Parqueadero"), ("wifi", "Wifi"),
                            ("latitud", "Ubicación (latitud/longitud)")]:
        if not info.get(clave):
            rv.add("Alta" if clave in ("direccion", "horario_general") else "Media", "Información general", etiqueta,
                   "No está cargado.", f"Cárgalo en Información general del panel ({etiqueta}).")
    if info.get("parqueadero") and len(str(info["parqueadero"])) < 25:
        rv.add("Media", "Información general", "Parqueadero",
               f"Solo dice «{info['parqueadero']}»: faltan tarifas, minutos gratis y mensualidad.",
               "Escribe el texto completo y correcto del parqueadero.")

    # Locales
    from models.store import Store
    tiendas = db.query(Store).all()
    nombres = Counter(_norm(s.name).strip() for s in tiendas if s.active)
    for s in tiendas:
        if not s.active:
            continue
        reg = s.name or f"(sin nombre, id {s.id})"
        n = _norm(s.name)
        if not s.name or len(s.name.strip()) < 3 or ":" in s.name or n.strip() in ("telefono", "telefonos", "local"):
            rv.add("Alta", "Locales", reg, "El nombre no es un nombre de local (parece un dato mal importado).",
                   "Corrige el nombre o desactiva el registro.", s.id)
        if nombres[n.strip()] > 1:
            rv.add("Media", "Locales", reg, "Hay otro local activo con el mismo nombre.", "Deja uno solo.", s.id)
        tel = re.sub(r"\D", "", s.phone or "")
        if not tel:
            rv.add("Media", "Locales", reg, "No tiene teléfono / WhatsApp.",
                   "Sin número, Any no puede dar contacto ni hacer domicilios.", s.id)
        elif not re.search(r"3\d{9}", tel):
            rv.add("Baja", "Locales", reg, f"El teléfono «{s.phone}» no parece un celular (WhatsApp).",
                   "Para el link de WhatsApp hace falta un celular de 10 dígitos.", s.id)
        zona = (s.location_hint or "").strip()
        es_torre = "torre" in _norm(f"{s.category} {s.location_hint} {s.description}")
        if not zona and not es_torre:
            rv.add("Alta", "Locales", reg, "No tiene zona.", "Pon la zona (Pasillo, Plaza, Terrazas…).", s.id)
        elif zona and not es_torre and not any(_norm(z) in _norm(zona) for z in ZONAS_OFICIALES):
            rv.add("Media", "Locales", reg, f"Zona «{zona}» no es una de las zonas oficiales.",
                   "Usa el nombre exacto de la zona.", s.id)
        if re.search(r"\b(piso|nivel)\s*\d", _norm(zona)) and not es_torre:
            rv.add("Media", "Locales", reg, f"La zona menciona piso o nivel: «{zona}».",
                   "Any ubica por zonas, no por pisos.", s.id)
        if es_torre and (s.floor or "").lower().startswith("por confirmar"):
            rv.add("Baja", "Torre Médica", reg, "El piso dice «Por confirmar».", "Carga el piso real.", s.id)
        if not (s.description or "").strip() and not (s.tags or "").strip():
            rv.add("Alta", "Locales", reg, "No tiene descripción ni productos.",
                   "Sin esto Any no lo encuentra cuando buscan lo que vende.", s.id)
        if not (s.schedule or "").strip():
            rv.add("Baja", "Locales", reg, "No tiene horario.", "Carga el horario del local.", s.id)
        cat = _norm(s.category)
        if "moda" in cat and re.search(r"joyer|bisuter|accesorio|artesan|libr|jugueter", n):
            rv.add("Media", "Locales", reg, f"Su rubro es «{s.category}» pero el nombre indica otro negocio.",
                   "Revisa el rubro: por eso aparece cuando preguntan por ropa.", s.id)
        if "licor" in cat and re.search(r"jumbo|supermerc|mercado", n + " " + _norm(s.description)):
            rv.add("Media", "Locales", reg, f"Su rubro es «{s.category}» pero es un supermercado.",
                   "Revisa el rubro.", s.id)
        if re.search(r"restaur|comida|plazoleta", cat) and not (s.catalogo_link or "").strip():
            rv.add("Baja", "Locales", reg, "Restaurante sin carta (ni link ni PDF).",
                   "Si tiene carta, súbela: Any la envía cuando la piden.", s.id)

    # Eventos
    from models.event import Event
    for e in db.query(Event).all():
        reg = e.name or f"evento {e.id}"
        v = _vencido(e.date, hoy)
        if v:
            rv.add("Baja", "Eventos", reg, f"Ya pasó ({e.date}); Any no lo muestra.", "Puedes borrarlo.", e.id)
        if not e.date:
            rv.add("Alta", "Eventos", reg, "No tiene fecha.", "Carga la fecha (dd/mm/aaaa).", e.id)
        if not e.location:
            rv.add("Media", "Eventos", reg, "No tiene lugar.", "Carga el lugar.", e.id)
        if not e.time and not re.search(r"domingo|sabado|lunes|todos", _norm(e.date)):
            rv.add("Baja", "Eventos", reg, "No tiene hora.", "Carga la hora.", e.id)
        pide_inscripcion = re.search(r"inscri|cupo|\$|precio|kit", _norm(e.description))
        if pide_inscripcion and not re.search(r"whatsapp|link|http|\d{7,}|mismo dia|en el lugar", _norm(e.description)):
            rv.add("Baja", "Eventos", reg, "La descripción no dice cómo inscribirse.",
                   "Any manda a la línea del centro comercial; si hay otro canal, escríbelo.", e.id)

    # Sorteos
    from models.raffle import Raffle
    for r in db.query(Raffle).all():
        reg = r.name or f"sorteo {r.id}"
        if r.active and _vencido(r.end_date, hoy):
            rv.add("Media", "Sorteos", reg, f"Está activo pero ya venció ({r.end_date}).", "Desactívalo.", r.id)
        if not r.end_date:
            rv.add("Media", "Sorteos", reg, "No tiene fecha de fin.", "Cárgala.", r.id)
        if _gritado(r.prize or "") or _gritado(r.description or ""):
            rv.add("Baja", "Sorteos", reg, "Texto escrito en MAYÚSCULAS.", "Escríbelo normal (Any lo suaviza).", r.id)
        if (r.prize or "").lower().startswith(("por cada", "cada ")):
            rv.add("Media", "Sorteos", reg, "En «premio» está la mecánica, no el premio.",
                   "Pon el premio en «premio» y la mecánica en la descripción.", r.id)

    # Promociones
    from models.marketing import Marketing
    for m in db.query(Marketing).all():
        reg = m.title or f"promoción {m.id}"
        if m.active and _vencido(m.valid_until, hoy):
            rv.add("Media", "Promociones", reg, f"Activa pero vencida ({m.valid_until}).", "Desactívala.", m.id)
        if not m.valid_until:
            rv.add("Baja", "Promociones", reg, "No tiene fecha de vigencia.", "Cárgala.", m.id)
        if len((m.description or "").strip()) < 60:
            rv.add("Media", "Promociones", reg, "La descripción es muy corta: no dice cómo se aplica ni en qué locales.",
                   "Escribe a quién aplica, en qué locales y cómo se usa.", m.id)
        if _gritado(m.description or ""):
            rv.add("Baja", "Promociones", reg, "Texto escrito en MAYÚSCULAS.", "Escríbelo normal.", m.id)

    # Conocimiento: contradicciones y temas que faltan
    from models.knowledge import KnowledgeEntry
    entradas = [k for k in db.query(KnowledgeEntry).all() if k.active]
    por_tema = defaultdict(list)
    for k in entradas:
        c = k.content or ""
        def _tarifas(patron):   # solo tarifas por hora (menos de $20.000); la mensualidad no cuenta
            out = set()
            for x in re.findall(patron, c, re.I):
                v = re.sub(r"\D", "", x)
                if v and int(v) < 20000:
                    out.add(v)
            return out
        precios_carro = _tarifas(r"carros?\s*(?:particular)?\s*:?\s*\$\s?([\d.,]+)")
        precios_moto = _tarifas(r"motos?\s*:?\s*\$\s?([\d.,]+)")
        if len(precios_carro) > 1 or len(precios_moto) > 1:
            rv.add("Alta", "Conocimiento", k.title,
                   f"Se contradice: tarifas de carro {sorted(precios_carro)} / moto {sorted(precios_moto)}.",
                   "Deja un solo valor. Ojo: si viene de la web (Fuente: elpuente.com.co) se vuelve a "
                   "descargar; quítale esa marca al editarla para que no se sobrescriba.", k.id)
        horarios = re.findall(r"\d{1,2}(?::\d{2})?\s*(?:am|pm|a\.m\.|p\.m\.)", c, re.I)
        if "horario" in _norm(c) and len(set(h.lower().replace(" ", "") for h in horarios)) > 4:
            rv.add("Media", "Conocimiento", k.title, "Trae varios horarios distintos; revisa que no se contradigan.",
                   "Deja un solo horario correcto.", k.id)
        if re.search(r"\b(\d+(?:er|do|ro|to|º|°)?\s*piso|primer piso|segundo piso)\b", _norm(c)):
            rv.add("Baja", "Conocimiento", k.title, "Menciona pisos.", "Any no habla de pisos: usa zonas.", k.id)
        if "elpuente.com.co/servicios" in c:
            rv.add("Info", "Conocimiento", k.title, "Viene de la web del centro comercial y se actualiza solo cada día.",
                   "Si la corriges a mano, borra la marca «(Fuente: …)» para que no se sobrescriba.", k.id)
        for tema, palabras in TEMAS.items():
            if any(p in _norm(k.title + " " + c) for p in palabras):
                por_tema[tema].append(k.title)
    if info.get("parqueadero"):
        por_tema["Parqueadero"].append("Información general")
    if info.get("wifi"):
        por_tema["Wifi"].append("Información general")
    for tema, palabras in TEMAS.items():
        if not por_tema.get(tema):
            rv.add("Alta" if tema in ("Wifi", "Baños", "Cajeros") else "Media", "Faltan datos", tema,
                   "No hay ninguna entrada que hable de esto: Any responde «no tengo información».",
                   f"Crea una entrada de conocimiento sobre {tema.lower()}.")

    # Cine
    try:
        from any.conocimiento import cine_datos as cd
        if not cd.cartelera(db):
            rv.add("Alta", "Cine", "Cartelera", "No hay funciones activas.", "Sincroniza o carga la cartelera.")
        if not cd.tarifas(db):
            rv.add("Media", "Cine", "Tarifas", "No se encontraron las tarifas en el local del cine.",
                   "Revisa el bloque «TARIFAS» del cine.")
        if not cd.confiteria(db):
            rv.add("Baja", "Cine", "Confitería", "No hay combos cargados.", "Revisa el bloque «CONFITERÍA».")
    except Exception as e:
        rv.add("Info", "Cine", "Revisión", f"No se pudo revisar: {e}", "")

    orden = {"Alta": 0, "Media": 1, "Baja": 2, "Info": 3}
    return sorted(rv.filas, key=lambda f: (orden.get(f[0], 9), f[1], str(f[2])))


TEMAS = {
    "Wifi": ["wifi", "wi-fi", "internet"],
    "Baños": ["bano", "sanitario"],
    "Cajeros": ["cajero", "atm"],
    "Parqueadero": ["parqueadero", "parqueo"],
    "Mascotas": ["mascota", "pet friendly", "perro"],
    "Punto de Información": ["punto de informacion"],
    "Tarjetas regalo": ["bono", "tarjeta regalo", "tarjeta de regalo"],
    "Auditorio y salones": ["auditorio", "salon de eventos", "salones"],
    "Enfermería / primeros auxilios": ["enfermeria", "primeros auxilios"],
    "Objetos perdidos": ["objetos perdidos", "perdi", "extravi"],
}


# ── vista de Any ──────────────────────────────────────────────────

def vista_any(db: Session) -> tuple[list[str], list[list]]:
    from any.conocimiento.catalogo import locales
    cab = ["id", "nombre", "zona", "rubro", "qué vende (lo que usa la búsqueda)", "WhatsApp", "teléfono", "horario",
           "carta (archivo)", "carta (link)", "Instagram", "fotos", "Torre Médica", "servicio interno"]
    filas = []
    for l in locales(db):
        ig = None
        try:
            from any.flujos.locales import instagram_de
            ig = instagram_de(l)
        except Exception:
            pass
        filas.append([l.id, l.nombre, l.zona, l.rubro, ", ".join(l.productos), l.whatsapp or "", l.telefono or "",
                      l.horario or "", l.fotos.get("carta") or "", l.carta_link or "", ig or "",
                      ", ".join(f"{k}: {v}" for k, v in l.fotos.items()), "sí" if l.es_torre else "",
                      "sí" if l.es_servicio else ""])
    return cab, filas


def config_any() -> list[tuple[str, list[str], list[list]]]:
    base = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "any")
    hojas = []
    import yaml

    def aplanar(prefijo, x, out):
        if isinstance(x, dict):
            for k, v in x.items():
                aplanar(f"{prefijo}.{k}" if prefijo else str(k), v, out)
        elif isinstance(x, list) and x and all(not isinstance(v, (dict, list)) for v in x):
            out.append([prefijo, ", ".join(map(str, x))])
        elif isinstance(x, list):
            for i, v in enumerate(x):
                aplanar(f"{prefijo}[{i}]", v, out)
        else:
            out.append([prefijo, x])
    for nombre, archivo in [("Any · mensajes", "mensajes.yaml"), ("Any · sinónimos", "sinonimos.yaml"),
                            ("Any · reglas", "reglas.yaml")]:
        filas = []
        try:
            with open(os.path.join(base, "datos", archivo), encoding="utf-8") as f:
                aplanar("", yaml.safe_load(f) or {}, filas)
        except Exception as e:
            filas = [["error", str(e)]]
        hojas.append((nombre, ["clave", "valor"], filas))
    prompts = []
    carpeta = os.path.join(base, "prompts")
    for raiz, _, archivos in os.walk(carpeta):
        for a in sorted(archivos):
            if a.endswith(".md"):
                ruta = os.path.join(raiz, a)
                with open(ruta, encoding="utf-8") as f:
                    prompts.append([os.path.relpath(ruta, carpeta), f.read()])
    hojas.append(("Any · prompts", ["archivo", "contenido"], prompts))
    try:
        from any.config import cfg
        c = cfg()
        filas = [["modelo del router", c.llm_router], ["modelo de los temas", c.llm_flujos],
                 ["modelo de respaldo", c.llm_respaldo or "ninguno"],
                 ["Any 2.0 responde WhatsApp (ANY_PRINCIPAL)", os.environ.get("ANY_PRINCIPAL", "")],
                 ["solo números de prueba (ANY_NUMEROS_PRUEBA)",
                  "sí" if os.environ.get("ANY_NUMEROS_PRUEBA") else "no (todos los clientes)"]]
    except Exception as e:
        filas = [["error", str(e)]]
    hojas.append(("Any · configuración", ["parámetro", "valor"], filas))
    return hojas


# ── armado del Excel ──────────────────────────────────────────────

def construir_excel(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    usados: set = set()
    negrita = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0F6E56")

    def hoja(titulo, cabecera, filas, anchos=None):
        ws = wb.create_sheet(_nombre_hoja(titulo, usados))
        ws.append([_celda(c) for c in cabecera])
        for c in ws[1]:
            c.font, c.fill = negrita, fondo
        for f in filas[:MAX_FILAS]:
            ws.append([_celda(v) for v in f])
        ws.freeze_panes = "A2"
        if cabecera:
            ws.auto_filter.ref = ws.dimensions
        for i, ancho in enumerate(anchos or [], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = ancho
        return ws

    insp = inspect(db.get_bind())
    tablas = sorted(insp.get_table_names())
    conteos = {}
    for t in tablas:
        try:
            conteos[t] = db.execute(text(f'SELECT COUNT(*) FROM "{t}"')).scalar()
        except Exception:
            db.rollback()
            conteos[t] = "error"

    # 1. Léeme
    ws = wb.active
    ws.title = _nombre_hoja("Léeme", usados)
    ws.append(["Respaldo completo de Any 2.0 y del panel — Centro Comercial El Puente"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generado: {datetime.now():%Y-%m-%d %H:%M}"])
    ws.append(["Contiene datos personales de clientes (teléfonos, conversaciones). Guárdalo con cuidado."])
    ws.append([])
    ws.append(["Hoja", "Qué tiene"])
    for a, b in [("Revisión", "Lo que falta o está mal en los datos, ordenado por gravedad. Empieza por aquí."),
                 ("Así ve Any los locales", "Cada local tal como lo usa Any 2.0."),
                 ("Any · …", "Mensajes fijos, sinónimos, reglas de promociones, prompts y modelos de Any 2.0."),
                 ("Una hoja por tabla", "Todas las tablas de la base de datos con todas sus columnas.")]:
        ws.append([a, b])
    ws.append([])
    ws.append(["Tabla", "Registros"])
    for t in tablas:
        ws.append([t, conteos[t]])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 90

    # 2. Revisión
    filas_rev = revisar(db)
    hoja("Revisión", ["gravedad", "área", "registro", "problema", "qué hacer", "id"], filas_rev,
         [9, 18, 34, 70, 70, 8])

    # 3. Vista de Any
    try:
        cab, filas = vista_any(db)
        hoja("Así ve Any los locales", cab, filas, [6, 30, 26, 24, 60, 30, 14, 30, 30, 30, 30, 40, 8, 8])
    except Exception as e:
        hoja("Así ve Any los locales", ["error"], [[str(e)]])

    # 4. Configuración de Any
    for titulo, cab, filas in config_any():
        hoja(titulo, cab, filas, [40, 120])

    # 5. Tablas completas
    for t in tablas:
        try:
            cols = [c["name"] for c in insp.get_columns(t)]
            visibles = [c for c in cols if c.lower() not in COLUMNAS_SECRETAS]
            lista = ", ".join(f'"{c}"' for c in visibles)
            orden = ' ORDER BY "id"' if "id" in cols else ""
            filas = db.execute(text(f'SELECT {lista} FROM "{t}"{orden}')).fetchall()
            hoja(t, visibles, [list(f) for f in filas])
        except Exception as e:
            db.rollback()
            hoja(t, ["error"], [[f"No se pudo leer: {e}"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/excel")
def respaldo_excel(db: Session = Depends(get_db)):
    contenido = construir_excel(db)
    nombre = f"respaldo_any_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    return StreamingResponse(io.BytesIO(contenido),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})