"""
📄 ARCHIVO: backend/routers/verificacion.py

Planilla de VERIFICACIÓN de datos: para revisar con el centro comercial, uno por uno, que todo
lo que sabe Any sea exacto, y aplicar las correcciones de una sola vez.

  GET  /verificacion/excel      descarga la planilla (una hoja por tipo de dato)
  POST /verificacion/importar   sube la planilla corregida: aplica SOLO lo que cambió, por id

Cómo se usa:
  1. Descargar la planilla. Cada fila es un registro (local, evento, sorteo, promoción,
     entrada de conocimiento, información general) con su id.
  2. Corregir directamente en las celdas blancas. Las grises son solo para leer.
     En "revisado" escribir "sí" cuando el registro quedó confirmado.
     Para apagar un registro, poner "no" en "activo".
  3. Subir el archivo. Solo se cambian las celdas distintas a lo que hay en la base.
     No se crean ni se borran registros: los nuevos se crean en el panel.

El Instagram de un local no tiene campo propio en la base: Any lo lee del link de la carta o de
la descripción. Si se cambia en la planilla, se escribe en la descripción ("Instagram: …").
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from models.database import get_db
from routers.auth import get_current_user

router = APIRouter(prefix="/verificacion", tags=["verificacion"], dependencies=[Depends(get_current_user)])

# hoja → (modelo, [(columna en la planilla, campo en la base o especial, editable, ancho)])
def _modelos():
    from models.event import Event
    from models.knowledge import KnowledgeEntry
    from models.mall_info import MallInfo
    from models.marketing import Marketing
    from models.raffle import Raffle
    from models.store import Store
    return {
        "Locales": (Store, [
            ("id", "id", False, 7), ("nombre", "name", True, 30), ("zona", "location_hint", True, 26),
            ("rubro", "category", True, 26), ("qué vende (palabras clave)", "tags", True, 45),
            ("descripción", "description", True, 60), ("horario", "schedule", True, 30),
            ("teléfono / WhatsApp", "phone", True, 16), ("carta o menú (link)", "catalogo_link", True, 35),
            ("instagram", "@instagram", True, 35), ("piso (solo Torre Médica)", "floor", True, 14),
            ("consultorio / número", "local_number", True, 14), ("activo", "active", True, 8),
            ("fotos cargadas", "@fotos", False, 22), ("problemas detectados", "@problemas", False, 45),
            ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
        "Eventos": (Event, [
            ("id", "id", False, 7), ("nombre", "name", True, 30), ("fecha (dd/mm/aaaa)", "date", True, 22),
            ("hora", "time", True, 16), ("lugar", "location", True, 28), ("descripción", "description", True, 60),
            ("prioridad (1-5)", "priority", True, 10), ("afiche", "photo_url", False, 30),
            ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
        "Sorteos": (Raffle, [
            ("id", "id", False, 7), ("nombre", "name", True, 30), ("premio", "prize", True, 35),
            ("cómo participar", "requirements", True, 45), ("fecha fin (dd/mm/aaaa)", "end_date", True, 18),
            ("dónde registrar", "location", True, 25), ("descripción", "description", True, 50),
            ("prioridad (1-5)", "priority", True, 10), ("activo", "active", True, 8),
            ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
        "Promociones": (Marketing, [
            ("id", "id", False, 7), ("título", "title", True, 30), ("descripción", "description", True, 60),
            ("local (id)", "store_id", True, 10), ("válido hasta (dd/mm/aaaa)", "valid_until", True, 18),
            ("prioridad (1-5)", "priority", True, 10), ("activo", "active", True, 8),
            ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
        "Conocimiento": (KnowledgeEntry, [
            ("id", "id", False, 7), ("título", "title", True, 40), ("contenido", "content", True, 100),
            ("activo", "active", True, 8), ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
        "Información general": (MallInfo, [
            ("id", "id", False, 7), ("nombre", "name", True, 28), ("dirección", "address", True, 35),
            ("horario general", "general_schedule", True, 70), ("teléfono", "phone", True, 16),
            ("parqueadero", "parking", True, 50), ("wifi", "wifi", True, 35), ("latitud", "latitude", True, 12),
            ("longitud", "longitude", True, 12), ("revisado", "@revisado", True, 10), ("notas", "@notas", True, 30)]),
    }


_IG = re.compile(r"(?:https?://)?(?:www\.)?instagram\.com/[A-Za-z0-9_.\-/?=&]+", re.I)


def _instagram(store) -> str:
    for t in (store.catalogo_link or "", store.description or ""):
        m = _IG.search(t)
        if m:
            u = m.group(0).rstrip(".,)").split("?")[0]
            return u if u.startswith("http") else "https://" + u
    return ""


def _sin_ig(t) -> str:
    t = _IG.sub("", str(t or ""))
    t = re.sub(r"(?i)\|?\s*instagram:\s*", " ", t)
    return re.sub(r"[\s|]+", " ", t).strip()


def _url_igual(u: str) -> str:
    """Para comparar links: sin https://, sin www., sin parámetros ni barra final."""
    u = (u or "").strip().lower().split("?")[0].rstrip("/")
    return re.sub(r"^https?://(www\.)?", "", u)


def _a_celda(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "sí" if v else "no"
    return v


def _a_valor(campo: str, v, actual):
    """Convierte lo escrito en la planilla al tipo del campo."""
    if v is None:
        v = ""
    if campo == "active":
        t = str(v).strip().lower()
        return t in ("sí", "si", "s", "1", "true", "x", "activo") if t else actual
    if campo in ("priority", "store_id"):
        t = str(v).strip()
        if not t:
            return None if campo == "store_id" else actual
        try:
            return int(float(t))
        except ValueError:
            raise ValueError(f"«{v}» no es un número")
    if campo in ("latitude", "longitude"):
        t = str(v).strip().replace(",", ".")
        return t or actual
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def construir(db: Session) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill

    problemas: dict = {}
    try:
        from routers.respaldo import revisar
        for g, area, reg, prob, hacer, id_ in revisar(db):
            if area in ("Locales", "Torre Médica") and id_:
                problemas.setdefault(int(id_), []).append(f"[{g}] {prob}")
    except Exception:
        pass
    fotos: dict = {}
    try:
        from models.store_photo import StorePhoto
        for f in db.query(StorePhoto).all():
            fotos.setdefault(f.store_id, []).append(f.label or "otra")
    except Exception:
        db.rollback()

    wb = Workbook()
    ws = wb.active
    ws.title = "Cómo usar"
    for fila in [
        ["Planilla de verificación — Any 2.0 · Centro Comercial El Puente"],
        [f"Generada: {datetime.now():%Y-%m-%d %H:%M}"], [],
        ["1. Revisa cada hoja fila por fila con el centro comercial (o con cada local)."],
        ["2. Corrige directamente en las celdas BLANCAS. Las grises (id, fotos, problemas) son solo para leer."],
        ["3. Cuando un registro quede confirmado, escribe «sí» en la columna «revisado»."],
        ["4. Para apagar un registro (duplicado, cerrado, mal importado) pon «no» en «activo»."],
        ["5. Sube el archivo en el panel con «Importar verificación». Solo se aplica lo que cambió."],
        [], ["No borres filas ni cambies los id. Los registros nuevos se crean desde el panel."],
        ["Zonas válidas: Pasillo Barichara, Pasillo Vélez, Pasillo Socorro, Pasillo Charalá, Pasillo Páramo, "
         "Pasillo Curití, Plaza Principal, Zona Terrazas, Plazoleta de Comidas La Ceiba, Paseo de la Cruz, "
         "Torre Médica, Oficinas Administrativas."],
        ["Fechas siempre en número: 31/10/2026. Teléfonos: el celular de 10 dígitos (para el link de WhatsApp)."],
    ]:
        ws.append(fila)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    gris = PatternFill("solid", fgColor="E8ECEA")
    cab_fill = PatternFill("solid", fgColor="0F6E56")
    alerta = PatternFill("solid", fgColor="FBE6DF")
    for hoja, (modelo, cols) in _modelos().items():
        sh = wb.create_sheet(hoja)
        sh.append([c[0] for c in cols])
        for i, (titulo, campo, editable, ancho) in enumerate(cols, start=1):
            celda = sh.cell(1, i)
            celda.font, celda.fill = Font(bold=True, color="FFFFFF"), cab_fill
            sh.column_dimensions[celda.column_letter].width = ancho
            if not editable:
                celda.comment = Comment("Solo lectura", "Any")
        q = db.query(modelo)
        filas = q.order_by(modelo.location_hint, modelo.name).all() if hoja == "Locales" else q.order_by(modelo.id).all()
        for obj in filas:
            fila = []
            for titulo, campo, editable, _ in cols:
                if campo == "@instagram":
                    fila.append(_instagram(obj))
                elif campo == "@fotos":
                    fila.append(", ".join(sorted(fotos.get(obj.id, []))) or "ninguna")
                elif campo == "@problemas":
                    fila.append("\n".join(problemas.get(obj.id, [])))
                elif campo.startswith("@"):
                    fila.append("")
                else:
                    fila.append(_a_celda(getattr(obj, campo, None)))
            sh.append(fila)
            r = sh.max_row
            for i, (titulo, campo, editable, _) in enumerate(cols, start=1):
                c = sh.cell(r, i)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if not editable:
                    c.fill = gris
                if campo == "@problemas" and c.value:
                    c.fill = alerta
        sh.freeze_panes = "C2"
        sh.auto_filter.ref = sh.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/excel")
def descargar(db: Session = Depends(get_db)):
    return StreamingResponse(io.BytesIO(construir(db)),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition":
                                      f'attachment; filename="verificacion_any_{datetime.now():%Y-%m-%d}.xlsx"'})


@router.post("/importar")
async def importar(file: UploadFile = File(...), db: Session = Depends(get_db)):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"No se pudo leer el Excel: {str(e)[:200]}")
    cambios, revisados, errores = [], 0, []
    for hoja, (modelo, cols) in _modelos().items():
        if hoja not in wb.sheetnames:
            continue
        sh = wb[hoja]
        filas = list(sh.iter_rows(values_only=True))
        if not filas:
            continue
        cab = [str(c or "").strip() for c in filas[0]]
        pos = {titulo: cab.index(titulo) for titulo, *_ in cols if titulo in cab}
        if "id" not in pos:
            errores.append(f"{hoja}: falta la columna id")
            continue
        for n, fila in enumerate(filas[1:], start=2):
            id_ = fila[pos["id"]] if pos["id"] < len(fila) else None
            if id_ in (None, ""):
                continue
            obj = db.get(modelo, int(id_))
            if not obj:
                errores.append(f"{hoja} fila {n}: no existe el id {id_}")
                continue
            for titulo, campo, editable, _ in cols:
                if not editable or titulo not in pos or pos[titulo] >= len(fila):
                    continue
                nuevo = fila[pos[titulo]]
                if campo == "@revisado":
                    revisados += str(nuevo or "").strip().lower() in ("sí", "si", "x", "ok")
                    continue
                if campo == "@notas":
                    continue
                if campo == "@instagram":
                    nuevo_ig = str(nuevo or "").strip()
                    if nuevo_ig and _url_igual(nuevo_ig) != _url_igual(_instagram(obj) or ""):
                        desc = _IG.sub("", obj.description or "").strip()
                        desc = re.sub(r"(?i)\s*\|?\s*instagram:\s*$", "", desc)
                        obj.description = (desc + f" | Instagram: {nuevo_ig}").strip(" |")
                        cambios.append(f"{hoja} · {obj.id} · instagram → {nuevo_ig}")
                    continue
                actual = getattr(obj, campo, None)
                if campo == "description" and _sin_ig(actual) == _sin_ig(nuevo):
                    continue      # solo difiere en el Instagram, que se maneja en su propia columna
                try:
                    valor = _a_valor(campo, nuevo, actual)
                except ValueError as e:
                    errores.append(f"{hoja} fila {n} ({titulo}): {e}")
                    continue
                if campo in ("name", "title", "content", "category", "floor") and not valor:
                    continue      # un campo obligatorio no se deja vacío
                if (actual if actual is not None else "") != (valor if valor is not None else "") and \
                        str(actual or "").strip() != str(valor or "").strip():
                    setattr(obj, campo, valor)
                    cambios.append(f"{hoja} · {obj.id} · {titulo}: «{str(_a_celda(actual))[:40]}» → «{str(_a_celda(valor))[:40]}»")
    if errores and not cambios:
        db.rollback()
        return {"ok": False, "cambios": 0, "errores": errores[:50]}
    db.commit()
    try:
        from routers.api import _reindex
        _reindex(db)
    except Exception:
        pass
    try:
        from any.conocimiento import catalogo
        catalogo.invalidar()
    except Exception:
        pass
    return {"ok": True, "cambios": len(cambios), "detalle": cambios[:300], "revisados": revisados,
            "errores": errores[:50]}