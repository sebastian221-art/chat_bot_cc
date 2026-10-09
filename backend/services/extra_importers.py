# 📄 ARCHIVO: backend/services/extra_importers.py
"""
Importadores de las plantillas NUEVAS del Centro Comercial El Puente:

  - parse_administrativo(bytes)  → la plantilla "Administrativo":
        * 01_Informacion_Institucional  → servicios del mall (Administración,
          Punto de Información, Parqueaderos)
        * 02_Eventos_y_Concursos        → se SEPARA en Eventos y Sorteos
          según la columna "Tipo de Actividad"
        * 03_Promociones_y_Banners      → promociones (Marketing)

  - parse_cine(bytes)            → la plantilla "Cine":
        * 01_Cartelera_Semanal + 02_Horarios_Funciones → funciones de cine
          (se cruzan por título)
        * 03_Tarifas_y_Promociones + 04_Confiteria_y_Combos → texto de
          tarifas/confitería que se guarda en la info del Cine para que Any
          responda precios

Ambos son TOLERANTES: si falta una hoja o una columna, no revientan;
devuelven lo que sí pudieron leer. La prioridad texto (Baja/Media/Alta) se
convierte al número que usa el sistema (1–5). Las fechas de Excel se
convierten a texto legible.
"""
import io
import datetime
import logging

import openpyxl

logger = logging.getLogger("mall_bot")


# ── utilidades ─────────────────────────────────────────────────────
def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def _headers(ws) -> dict:
    h = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val:
            h[str(val).lower().strip()] = c
    return h


def _col(headers, *nombres):
    for nombre in nombres:
        n = nombre.lower()
        for clave, col in headers.items():
            if n in clave:
                return col
    return None


def _prioridad_a_numero(texto: str, default: int = 3) -> int:
    """Baja→2, Media→3, Alta→5 (lo que usa el sistema: 1–5)."""
    t = (texto or "").strip().lower()
    if t.startswith("baj"):
        return 2
    if t.startswith("med"):
        return 3
    if t.startswith("alt"):
        return 5
    return default


def _rango_fecha(ini: str, fin: str) -> str:
    """Devuelve 'DD/MM/AAAA' o 'DD/MM/AAAA al DD/MM/AAAA' si hay rango."""
    ini, fin = (ini or "").strip(), (fin or "").strip()
    if ini and fin and ini != fin:
        return f"{ini} al {fin}"
    return ini or fin


def _hoja(wb, *nombres):
    """Busca una hoja por coincidencia parcial de nombre."""
    for n in nombres:
        for sh in wb.sheetnames:
            if n.lower() in sh.lower():
                return wb[sh]
    return None


# ── ADMINISTRATIVO ─────────────────────────────────────────────────
def parse_administrativo(file_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    servicios, eventos, sorteos, promos = [], [], [], []

    # ── Hoja 1: Información Institucional → servicios del mall ──────
    ws = _hoja(wb, "informacion_institucional", "institucional")
    if ws:
        h = _headers(ws)
        c_area = _col(h, "área", "area", "servicio")
        c_ubic = _col(h, "ubicación", "ubicacion")
        c_hor  = _col(h, "horario")
        c_tel  = _col(h, "teléfono", "telefono", "whatsapp")
        c_mail = _col(h, "correo", "email")
        c_info = _col(h, "información", "informacion", "procedimiento")
        for row in range(2, ws.max_row + 1):
            area = _txt(ws.cell(row, c_area).value) if c_area else ""
            if not area:
                continue
            ubic = _txt(ws.cell(row, c_ubic).value) if c_ubic else ""
            hor  = _txt(ws.cell(row, c_hor).value) if c_hor else ""
            tel  = _txt(ws.cell(row, c_tel).value) if c_tel else ""
            mail = _txt(ws.cell(row, c_mail).value) if c_mail else ""
            info = _txt(ws.cell(row, c_info).value) if c_info else ""
            # Primer número como principal; TODOS los números/correos
            # (con su etiqueta de área) van a la descripción para que Any
            # pueda dar el de un área específica si se lo piden.
            import re as _re
            nums = [n for n in _re.sub(r"[^\d]+", " ", tel).split() if len(n) >= 7]
            principal = nums[0] if nums else _re.sub(r"[^\d]", "", tel)
            desc_partes = []
            if info:
                desc_partes.append(info)
            if tel:
                desc_partes.append(f"Contactos: {tel}")
            if mail:
                desc_partes.append(f"Correos: {mail}")
            servicios.append({
                "name": area,
                "location_hint": ubic,
                "schedule": hor,
                "phone": principal[:60],
                "description": ". ".join(p for p in desc_partes if p),
                "category": "Servicios del Centro Comercial",
            })

    # ── Hoja 2: Eventos y Concursos → se separa en eventos/sorteos ──
    ws = _hoja(wb, "eventos", "concursos")
    if ws:
        h = _headers(ws)
        c_nom  = _col(h, "nombre del evento", "nombre")
        c_tipo = _col(h, "tipo de actividad", "tipo")
        c_ini  = _col(h, "fecha de inicio", "inicio")
        c_fin  = _col(h, "fecha de fin", "fin")
        c_lug  = _col(h, "lugar", "pasillo")
        c_hor  = _col(h, "horario")
        c_desc = _col(h, "descripción", "descripcion", "mecánica", "mecanica")
        c_loc  = _col(h, "locales involucrados", "involucrados")
        for row in range(2, ws.max_row + 1):
            nom = _txt(ws.cell(row, c_nom).value) if c_nom else ""
            if not nom:
                continue
            tipo = (_txt(ws.cell(row, c_tipo).value) if c_tipo else "").lower()
            ini  = _txt(ws.cell(row, c_ini).value) if c_ini else ""
            fin  = _txt(ws.cell(row, c_fin).value) if c_fin else ""
            lug  = _txt(ws.cell(row, c_lug).value) if c_lug else ""
            hor  = _txt(ws.cell(row, c_hor).value) if c_hor else ""
            desc = _txt(ws.cell(row, c_desc).value) if c_desc else ""
            loc  = _txt(ws.cell(row, c_loc).value) if c_loc else ""
            fecha = _rango_fecha(ini, fin)
            desc_full = desc + (f" | Locales: {loc}" if loc else "")

            es_sorteo = "sorteo" in tipo or "concurso" in tipo
            if es_sorteo:
                prio = 5 if "principal" in nom.lower() else 4
                sorteos.append({
                    "name": nom,
                    "prize": (desc[:240] if desc else nom)[:250],  # el premio va descrito en la mecánica
                    "requirements": desc_full,
                    "end_date": fin or fecha,
                    "location": lug,
                    "description": desc_full,
                    "priority": prio,
                })
            else:
                # El campo 'time' solo admite 30 caracteres. Si el horario
                # es largo (ej. "Viernes 7-10 PM - Sáb y Dom 3-10 PM") va a
                # la descripción y 'time' queda corto o vacío, para no
                # reventar la base (error "value too long").
                time_corto = hor if len(hor) <= 30 else ""
                desc_evt = desc_full
                if hor and not time_corto:
                    desc_evt = (desc_evt + " | Horario: " + hor).strip(" |")
                eventos.append({
                    "name": nom,
                    "date": fecha,
                    "time": time_corto,
                    "location": lug,
                    "description": desc_evt,
                    "priority": 3,
                })

    # ── Hoja 3: Promociones y Banners → Marketing ──────────────────
    ws = _hoja(wb, "promociones", "banners")
    if ws:
        h = _headers(ws)
        c_tit  = _col(h, "título de promoción", "titulo de promoción", "título", "titulo")
        c_prio = _col(h, "prioridad")
        c_loc  = _col(h, "locales", "sectores")
        c_exp  = _col(h, "expiración", "expiracion", "vence")
        c_txt  = _col(h, "texto de cierre", "whatsapp", "texto")
        for row in range(2, ws.max_row + 1):
            tit = _txt(ws.cell(row, c_tit).value) if c_tit else ""
            if not tit:
                continue
            prio = _prioridad_a_numero(_txt(ws.cell(row, c_prio).value) if c_prio else "")
            loc  = _txt(ws.cell(row, c_loc).value) if c_loc else ""
            exp  = _txt(ws.cell(row, c_exp).value) if c_exp else ""
            txt  = _txt(ws.cell(row, c_txt).value) if c_txt else ""
            desc_partes = []
            if txt:
                desc_partes.append(txt)
            if loc:
                desc_partes.append(f"Aplica en: {loc}")
            promos.append({
                "title": tit,
                "description": ". ".join(desc_partes),
                "priority": prio,
                "valid_until": exp,
                "active": True,   # como pidió el cliente, las vencidas también entran
            })

    # ── RED DE SEGURIDAD: recortar a los límites reales de cada tabla ──
    # (Event/Raffle/Marketing no tienen auto-recorte como Store, así que
    # lo hacemos aquí para que nunca salga "value too long".)
    def _cap(lst, limites):
        for d in lst:
            for campo, lim in limites.items():
                v = d.get(campo)
                if v and len(str(v)) > lim:
                    d[campo] = str(v)[:lim]
    _cap(eventos, {"name": 150, "date": 30, "time": 30, "location": 150})
    _cap(sorteos, {"name": 150, "prize": 250, "end_date": 30, "location": 150})
    _cap(promos,  {"title": 150, "valid_until": 30})
    _cap(servicios, {"name": 150, "schedule": 200, "phone": 60, "location_hint": 200})

    return {"servicios": servicios, "eventos": eventos, "sorteos": sorteos, "promos": promos}


# ── CINE ───────────────────────────────────────────────────────────
def parse_cine(file_bytes: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    # ── Horarios: {titulo_normalizado: showtimes} ──────────────────
    horarios = {}
    ws = _hoja(wb, "horarios")
    if ws:
        h = _headers(ws)
        c_tit = _col(h, "título", "titulo", "película", "pelicula")
        c_hor = _col(h, "horarios de función", "horarios", "función", "funcion")
        for row in range(2, ws.max_row + 1):
            tit = _txt(ws.cell(row, c_tit).value) if c_tit else ""
            if not tit:
                continue
            hor = _txt(ws.cell(row, c_hor).value) if c_hor else ""
            horarios[tit.strip().lower()] = hor

    # ── Cartelera → funciones ──────────────────────────────────────
    funciones = []
    ws = _hoja(wb, "cartelera")
    if ws:
        h = _headers(ws)
        c_tit  = _col(h, "título", "titulo", "película", "pelicula")
        c_fmt  = _col(h, "formato")
        c_idi  = _col(h, "idioma", "subtítulos", "subtitulos")
        c_clas = _col(h, "clasificación", "clasificacion")
        c_gen  = _col(h, "género", "genero")
        c_dur  = _col(h, "duración", "duracion")
        c_sin  = _col(h, "sinopsis")
        c_tra  = _col(h, "tráiler", "trailer", "youtube")
        for row in range(2, ws.max_row + 1):
            tit = _txt(ws.cell(row, c_tit).value) if c_tit else ""
            if not tit:
                continue
            fmt  = _txt(ws.cell(row, c_fmt).value) if c_fmt else ""
            idi  = _txt(ws.cell(row, c_idi).value) if c_idi else ""
            clas = _txt(ws.cell(row, c_clas).value) if c_clas else ""
            gen  = _txt(ws.cell(row, c_gen).value) if c_gen else ""
            dur  = _txt(ws.cell(row, c_dur).value) if c_dur else ""
            sin  = _txt(ws.cell(row, c_sin).value) if c_sin else ""
            tra  = _txt(ws.cell(row, c_tra).value) if c_tra else ""
            # Descripción = ficha técnica + sinopsis + tráiler
            ficha = " · ".join(p for p in [gen, clas, fmt, idi, (dur if dur else "")] if p)
            desc_partes = []
            if ficha:
                desc_partes.append(ficha)
            if sin:
                desc_partes.append(sin)
            if tra:
                desc_partes.append(f"Tráiler: {tra}")
            showtimes = horarios.get(tit.strip().lower(), "")
            funciones.append({
                "title": tit[:150],
                "showtimes": showtimes[:200],
                "description": "\n".join(desc_partes),
                "is_premiere": False,
            })

    # ── Tarifas + Confitería → texto para la info del Cine ─────────
    bloques = []
    ws = _hoja(wb, "tarifas")
    if ws:
        h = _headers(ws)
        c_dia = _col(h, "día", "dia", "condición", "condicion")
        c_fmt = _col(h, "formato")
        c_pag = _col(h, "tipo de pago", "pago")
        c_sil = _col(h, "tipo de silla", "silla")
        c_pre = _col(h, "precio")
        c_obs = _col(h, "observaciones", "condiciones")
        lineas = []
        for row in range(2, ws.max_row + 1):
            dia = _txt(ws.cell(row, c_dia).value) if c_dia else ""
            if not dia:
                continue
            fmt = _txt(ws.cell(row, c_fmt).value) if c_fmt else ""
            pag = _txt(ws.cell(row, c_pag).value) if c_pag else ""
            sil = _txt(ws.cell(row, c_sil).value) if c_sil else ""
            pre = _txt(ws.cell(row, c_pre).value) if c_pre else ""
            obs = _txt(ws.cell(row, c_obs).value) if c_obs else ""
            lineas.append(f"  • {dia} — {fmt} {sil} ({pag}): ${pre}" + (f" [{obs}]" if obs else ""))
        if lineas:
            bloques.append("💵 TARIFAS DE CINE:\n" + "\n".join(lineas))

    ws = _hoja(wb, "confiteria", "confitería", "combos")
    if ws:
        h = _headers(ws)
        c_prod = _col(h, "producto", "combo")
        c_desc = _col(h, "descripción", "descripcion", "contenido")
        c_pre  = _col(h, "precio")
        c_pag  = _col(h, "tipo de pago", "pago")
        c_disp = _col(h, "disponibilidad")
        lineas = []
        for row in range(2, ws.max_row + 1):
            prod = _txt(ws.cell(row, c_prod).value) if c_prod else ""
            if not prod:
                continue
            d   = _txt(ws.cell(row, c_desc).value) if c_desc else ""
            pre = _txt(ws.cell(row, c_pre).value) if c_pre else ""
            pag = _txt(ws.cell(row, c_pag).value) if c_pag else ""
            lineas.append(f"  • {prod}: {d} — ${pre}" + (f" ({pag})" if pag else ""))
        if lineas:
            bloques.append("🍿 CONFITERÍA Y COMBOS:\n" + "\n".join(lineas))

    info_extra = "\n\n".join(bloques)
    return {"funciones": funciones, "info_extra": info_extra}